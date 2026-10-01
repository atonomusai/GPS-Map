# -*- coding: utf-8 -*-
"""Builds 'GPS Camera - APIs Explained Simply.docx'"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_SECTION

INDIGO = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_DARK = RGBColor(0x37, 0x2F, 0xC4)
DARK = RGBColor(0x1A, 0x1C, 0x20)
GRAY = RGBColor(0x5A, 0x63, 0x73)
LIGHT_BG = "EEF0FE"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# ---------- base page / font setup ----------
section = doc.sections[0]
section.page_width = Cm(21.59)
section.page_height = Cm(27.94)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.2


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_cell_borders(cell, color="D8DBFB", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def add_bottom_rule(paragraph, color="4F46E5", size="18"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def h_title(text, size=30, color=INDIGO, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    return p


def subtitle(text, size=14, color=GRAY, italic=True, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def h1(text, number=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(10)
    if number is not None:
        badge = p.add_run(f"  {number}  ")
        badge.font.size = Pt(13)
        badge.font.bold = True
        badge.font.color.rgb = WHITE
    r = p.add_run(("  " if number is not None else "") + text)
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = INDIGO
    add_bottom_rule(p)
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = INDIGO_DARK
    return p


def body(text, bold_lead=None, size=11.5, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.font.bold = True
        lead.font.size = Pt(size)
        lead.font.color.rgb = DARK
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = DARK
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.font.bold = True
        lead.font.size = Pt(11)
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


def callout(label, text, fill="F4F5FF", border="C7CBFA"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(16.3)
    cell = table.rows[0].cells[0]
    cell.width = Cm(16.3)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, color=border, sz="6")
    cell.margin_left = Cm(0.35)
    cell.paragraphs[0].paragraph_format.space_before = Pt(6)
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    lr = p0.add_run(label)
    lr.font.bold = True
    lr.font.size = Pt(11)
    lr.font.color.rgb = INDIGO_DARK
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_after = Pt(6)
    tr = p1.add_run(text)
    tr.font.size = Pt(11)
    tr.font.italic = True
    tr.font.color.rgb = DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def spacer(pt=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pt)
    p.paragraph_format.space_before = Pt(0)


def page_break():
    doc.add_page_break()


# =====================================================================
# COVER PAGE
# =====================================================================
for _ in range(4):
    doc.add_paragraph()

cover_icon = doc.add_paragraph()
cover_icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cover_icon.add_run("\U0001F4F1  \U0001F310  \U0001F4CD")
r.font.size = Pt(40)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_before = Pt(18)
r = t.add_run("The Helpers Behind")
r.font.size = Pt(30)
r.font.bold = True
r.font.color.rgb = DARK

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t2.add_run("GPS Camera")
r.font.size = Pt(40)
r.font.bold = True
r.font.color.rgb = INDIGO

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_before = Pt(14)
r = sub.add_run("A Simple, Friendly Guide to Every Outside Service\nthis App Talks To — and Why")
r.font.size = Pt(14)
r.font.italic = True
r.font.color.rgb = GRAY

for _ in range(6):
    doc.add_paragraph()

footer_note = doc.add_paragraph()
footer_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer_note.add_run("Written so that anyone — even a curious kid — can understand\nwhat's happening behind the scenes every time you open the app.")
r.font.size = Pt(11)
r.font.color.rgb = GRAY
r.font.italic = True

page_break()

# =====================================================================
# WHAT IS AN API
# =====================================================================
h_title("What Even Is an “API”?", size=22, space_after=10)

body("", bold_lead="")
p = doc.add_paragraph()
r = p.add_run(
    "Imagine you're playing outside and you want to know if it's going to rain. "
    "You don't own a weather satellite — so you call your mom, and she checks "
    "her phone and tells you: “Yes, bring an umbrella!”"
)
r.font.size = Pt(12)

p2 = doc.add_paragraph()
r = p2.add_run(
    "An API (say it: “A-P-I”) works exactly like that phone call. It's a way for one app "
    "to ask another app or service a question and get an answer back — instantly, "
    "automatically, without a human in the middle."
)
r.font.size = Pt(12)

callout(
    "\U0001F4A1  In one sentence:",
    "An API is like a helpful friend on the phone that an app can call anytime it "
    "needs an answer — like “where am I?”, “what street is this?”, or "
    "“what does this place look like right now?”"
)

spacer(4)
p3 = doc.add_paragraph()
r = p3.add_run(
    "GPS Camera is a photo and travel app, so it needs to ask a lot of these questions: "
    "Where is the user standing? What's this address called? How do I get from A to B? "
    "What does that live traffic camera see right now? It can't know any of that by itself — "
    "so it calls outside helpers (APIs) to find out, the same way you'd call a friend who "
    "knows the answer."
)
r.font.size = Pt(12)

spacer(6)
p4 = doc.add_paragraph()
r = p4.add_run("This document introduces every single helper GPS Camera calls on, one at a time.")
r.font.size = Pt(12)
r.font.bold = True

page_break()

# =====================================================================
# SECTION: MAP & LOCATION HELPERS INTRO
# =====================================================================
h_title("Meet the Helpers", size=22, space_after=4)
subtitle("There are two kinds: helpers that live on the internet, and helpers built right into your phone.", size=12, italic=True)
spacer(10)

# ---------------- API 1: OSM Map Tiles ----------------
h1("The Map Drawer", number="1")
h2("(OpenStreetMap Map Tiles)")

body(
    "Every screen in the app with a map on it — the Nearby Places map, the Route Finder, "
    "the location picker — needs actual map pictures to show you: streets, buildings, rivers, "
    "parks. GPS Camera doesn't draw those itself. It asks OpenStreetMap, a free, "
    "volunteer-built map of the entire world, to send over the little square picture "
    "“tiles” that get stitched together into the map you see and can zoom into.",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "A giant book of world maps that anyone can borrow pages from, for free, forever — "
    "built and kept up to date by regular people all over the world, like a "
    "Wikipedia but for maps."
)

bullet("Nearby Places screen (the map behind the search results)", bold_lead="Where you'll see it: ")
bullet("Route Finder (the road map with your route drawn on it)")
bullet("Pick-a-Location map screen, and the Satellite / Street View map screens")
bullet("Trip History maps (showing where you drove/walked)")

spacer(4)
body("Completely free, no sign-up needed — it's a shared public resource, kept alive by donations.", bold_lead="Cost: ")

page_break()

# ---------------- API 2: Nominatim Search ----------------
h1("The Address Finder", number="2")
h2("(OpenStreetMap Nominatim)")

body(
    "When you type a place name into a search box — like “Eiffel Tower” or “Main "
    "Street bakery” — something has to turn those words into an actual spot on the map "
    "(exact coordinates). This helper does exactly that job, both ways: name → coordinates, "
    "and coordinates → name (so the app can tell you “you are on Main Street” after it "
    "finds your GPS position).",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "A phone book, but instead of looking up a person's number, you look up a "
    "place's exact spot on Earth — or the other way around, you give it a spot "
    "and it tells you the place's name and address."
)

bullet("The search bar in the map / location picker screens", bold_lead="Where you'll see it: ")
bullet("Turning your GPS dot into a readable address like “Chishtian, Pakistan”")
bullet("The Custom Location and Nearby Places search boxes")

spacer(4)
body("Free to use — it's run by the same OpenStreetMap volunteer community.", bold_lead="Cost: ")

page_break()

# ---------------- API 3: OSRM ----------------
h1("The Direction Giver", number="3")
h2("(OSRM — Open Source Routing Machine)")

body(
    "This is the helper behind the Route Finder feature. You tell it where you are and "
    "where you want to go, and it works out the best path — by car, bike, or on foot — "
    "and hands back a line to draw on the map, plus how far it is and how long it'll take.",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "A friend who has memorized every road on Earth, and when you ask “how do "
    "I get from my house to the park?”, instantly draws you the best path on a "
    "napkin — with a time estimate."
)

bullet("Route Finder screen — the blue route line, the distance, and time estimate", bold_lead="Where you'll see it: ")
bullet("The “Start Navigation” flow when a route is chosen")

spacer(4)
body("Free, public routing service — no account or payment needed.", bold_lead="Cost: ")

page_break()

# ---------------- API 4: Overpass ----------------
h1("The Nearby-Things Finder", number="4")
h2("(Overpass API)")

body(
    "When you open Nearby Places and tap “Hospitals” or “Restaurants” or “Fuel "
    "Stations,” something needs to search a giant database of every tagged place "
    "on Earth and hand back the ones close to you. That's this helper's whole job — "
    "it searches OpenStreetMap's map data for things matching what you asked for, "
    "within a circle around your location.",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "Asking a local tour guide, “hey, is there a pharmacy near here?” — and "
    "they instantly point to every single one within a few streets, sorted by "
    "how close they are."
)

bullet("Nearby Places screen — every category chip (Hospitals, Cafes, Banks, Fuel, etc.)", bold_lead="Where you'll see it: ")
bullet("The list of results and their pins on the map")

spacer(4)
body("Free, public service. The app is polite about it — it tries several backup copies of this service if the main one is busy, so results still come back reliably.", bold_lead="Cost: ")

page_break()

# ---------------- API 5: Windy Webcams ----------------
h1("The Live Camera Viewer", number="5")
h2("(Windy Webcams API)")

body(
    "This powers the Live Street Camera feature — it lets you browse real public "
    "webcams from cities and landmarks all over the world, filter them by country "
    "or category (like Beach, Mountain, Traffic), and see the latest picture each "
    "camera has taken.",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "A wall of TV screens in a control room, each one showing a live snapshot "
    "from a real camera somewhere in the world — refreshed every so often, not "
    "a constant video, more like a postcard that updates itself every few "
    "minutes."
)

bullet("The “Live Street Camera” screen and its country / category filters", bold_lead="Where you'll see it: ")
bullet("The “Street View Cameras” row on the Home screen")
bullet("The camera detail screen, including the “Updated X minutes ago” note and the Timelapse player")

spacer(4)
body(
    "Free to use, with one small rule attached: the app must always show a little "
    "“Webcams provided by Windy.com” credit — like crediting a photographer whose "
    "picture you're using. You'll notice that note at the bottom of the camera screens.",
    bold_lead="Cost: "
)

page_break()

# =====================================================================
# SECTION: PHONE'S OWN BUILT-IN HELPERS
# =====================================================================
h_title("Helpers Built Right Into Your Phone", size=22, space_after=4)
subtitle("These aren't on the internet at all — they're services Android itself provides, using your phone's own hardware.", size=12, italic=True)
spacer(10)

# ---------------- API 6: Location/GPS ----------------
h1("The Position Sensor", number="6")
h2("(Android Location Service / GPS)")

body(
    "Your phone has a real GPS chip inside it — the same kind of technology that "
    "talks to satellites orbiting Earth. This helper is how the app asks the phone, "
    "“where exactly am I right now?” and gets back your precise latitude and "
    "longitude, plus how confident the phone is in that answer (its “accuracy”).",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "Playing a game of hot-and-cold with satellites in space — your phone "
    "listens to signals from several of them at once and does math to figure "
    "out exactly where you're standing on Earth, down to a few meters."
)

bullet("The GPS card on the Home screen (your city, address, and coordinates)", bold_lead="Where you'll see it: ")
bullet("The GPS stamp burned into every photo and video you capture")
bullet("Speedometer, Altimeter, Compass, and Trip Tracking — all of these need to know where you are and how fast you're moving")

spacer(4)
body("Free — it's simply using hardware already inside your phone. No internet needed once you have satellite signal.", bold_lead="Cost: ")

page_break()

# ---------------- API 7: Camera ----------------
h1("The Camera Controller", number="7")
h2("(Android CameraX)")

body(
    "This is how the app talks to your phone's actual camera lens — turning it on, "
    "showing you a live preview, adjusting focus and flash, and capturing the "
    "photo or video when you tap the shutter button.",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "A remote control for your camera lens — the app presses all the right "
    "buttons (focus, flash, capture) on your behalf, faster and more precisely "
    "than doing it by hand."
)

bullet("The Camera screen — live preview, taking photos, recording videos", bold_lead="Where you'll see it: ")
bullet("The Compass Camera and Import-and-Stamp screens")

spacer(4)
body("Free — built into every Android phone.", bold_lead="Cost: ")

page_break()

# ---------------- API 8: Sensors ----------------
h1("The Body-Sense Helpers", number="8")
h2("(Android Sensors: Compass, Accelerometer, Barometer)")

body(
    "Your phone has tiny sensors that can feel the world around it: one senses "
    "which way is North (like a tiny compass needle), one senses tilting and "
    "movement, and some phones even have one that senses air pressure to guess "
    "your height above sea level.",
    bold_lead="What it does: "
)

callout(
    "\U0001F9F8  Think of it like:",
    "A tiny set of superpowers hidden inside your phone: a built-in compass "
    "that always knows North, a balance sense that knows which way is “up”, "
    "and (on some phones) a weather-station-style sense for how high up you are."
)

bullet("The Compass screen (the spinning needle pointing North)", bold_lead="Where you'll see it: ")
bullet("The Altimeter & Elevation screen (your height and the barometer reading)")

spacer(4)
body("Free — built into the phone's hardware. On phones without a height sensor, the app calculates an estimate using GPS and math instead.", bold_lead="Cost: ")

page_break()

# =====================================================================
# QUICK REFERENCE TABLE
# =====================================================================
h_title("Quick Look: All Helpers at a Glance", size=20, space_after=14)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(4.0), Cm(5.3), Cm(5.0), Cm(2.0)]
for i, w in enumerate(widths):
    table.columns[i].width = w

hdr = table.rows[0].cells
headers = ["Helper", "What It's For", "Where You See It", "Free?"]
for i, htext in enumerate(headers):
    hdr[i].width = widths[i]
    set_cell_shading(hdr[i], "4F46E5")
    set_cell_borders(hdr[i], color="4F46E5", sz="4")
    p = hdr[i].paragraphs[0]
    r = p.add_run(htext)
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = WHITE

rows_data = [
    ("OpenStreetMap\n(Map Tiles)", "Drawing the maps you see on screen", "Every map screen", "Yes"),
    ("Nominatim\n(Search)", "Turning place names into map spots, and back again", "Search boxes, address lookup", "Yes"),
    ("OSRM\n(Routing)", "Working out directions between two points", "Route Finder", "Yes"),
    ("Overpass\n(Nearby Search)", "Finding hospitals, cafes, fuel stations etc. near you", "Nearby Places", "Yes"),
    ("Windy Webcams", "Showing live public camera snapshots worldwide", "Live Street Camera", "Yes (credit required)"),
    ("Android Location /\nGPS", "Finding your exact real-world position", "Home, Camera stamp, Speedometer, Trips", "Yes"),
    ("Android CameraX", "Controlling the phone's camera lens", "Camera screen", "Yes"),
    ("Android Sensors", "Sensing direction, tilt, and altitude", "Compass, Altimeter", "Yes"),
]

for row in rows_data:
    cells = table.add_row().cells
    for i, val in enumerate(row):
        cells[i].width = widths[i]
        set_cell_borders(cells[i], color="D8DBFB", sz="4")
        set_cell_shading(cells[i], "F7F8FF")
        p = cells[i].paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9.5)
        if i == 0:
            r.font.bold = True
            r.font.color.rgb = INDIGO_DARK

spacer(20)

closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = closing.add_run(
    "And that's it! Every time GPS Camera shows you a map, finds your location, "
    "snaps a photo, or shows a live street camera — it's really just politely "
    "asking one of these eight helpers a quick question, and passing the answer "
    "along to you."
)
r.font.size = Pt(12)
r.font.italic = True
r.font.color.rgb = GRAY

doc.save("GPS Camera - APIs Explained Simply.docx")
print("Saved.")
